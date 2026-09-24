import os
import json
import sqlite3
import requests
import threading
import time
from datetime import datetime
from flask import Flask, request, jsonify, render_template_string
from app.tools import web_search, read_file, write_file, ssh_exec, open_browser

app = Flask(__name__, static_folder='static', static_url_path='/static')
from app.file_upload import register_upload
register_upload(app)
OLLAMA_URL = "http://127.0.0.1:11434/api/chat"
MODEL = "qwen2.5:7b"
DB_PATH = "C:/AI_Assistant/memory.db"
MAX_AGENT_STEPS = 6   # جلوگیری از حلقه بی‌نهایت
MAX_HISTORY_MESSAGES = 20   # حداکثر تعداد پیام‌های نگه‌داشته‌شده از مکالمه (برای کنترل حجم)

# تاریخچه ساده مکالمه در حافظه (per-process). با ری‌استارت سرور پاک می‌شود.
conversation_history = []


# ---------- MEMORY & TASKS DATABASE SETUP ----------

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    # جدول حافظه
    cur.execute('''
        CREATE TABLE IF NOT EXISTS memories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            content TEXT,
            category TEXT
        )
    ''')
    # جدول وظایف زمان‌بندی‌شده (Task Scheduler)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            prompt TEXT NOT NULL,
            schedule_type TEXT NOT NULL, -- 'once', 'daily', 'interval'
            run_at TEXT,                 -- برای حالت once: 'YYYY-MM-DD HH:MM:SS'
            repeat_interval INTEGER,     -- برای حالت interval بر حسب دقیقه
            last_run TEXT,
            next_run TEXT,
            enabled INTEGER DEFAULT 1
        )
    ''')
    conn.commit()
    conn.close()

# اجرای مقداردهی اولیه دیتابیس هنگام بالا آمدن ماژول
init_db()


def get_memory():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("SELECT id, content, category FROM memories ORDER BY id DESC LIMIT 50")
    rows = cur.fetchall()
    conn.close()
    if not rows:
        return "حافظه خالی است."
    return "\n".join(f"{r[0]} [{r[2]}] {r[1]}" for r in rows)


def save_memory(text):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("INSERT INTO memories (content, category) VALUES (?, ?)", (text, "عمومی"))
    conn.commit()
    conn.close()


# ---------- TASK SCHEDULER (پس‌زمینه) ----------

class TaskScheduler(threading.Thread):
    def __init__(self, agent_runner_callback, interval=30):
        super().__init__()
        self.agent_runner_callback = agent_runner_callback
        self.interval = interval  # بررسی هر 30 ثانیه
        self.daemon = True
        self._stop_event = threading.Event()

    def run(self):
        print("⏰ TaskScheduler background thread started.")
        while not self._stop_event.is_set():
            try:
                self.check_and_run_tasks()
            except Exception as e:
                print(f"Error in TaskScheduler: {e}")
            time.sleep(self.interval)

    def check_and_run_tasks(self):
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        now_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        # استخراج تسک‌های فعال که زمان اجرایشان فرا رسیده است
        cursor.execute("SELECT * FROM tasks WHERE enabled = 1 AND next_run <= ?", (now_str,))
        tasks = cursor.fetchall()
        
        for task in tasks:
            task_id = task['id']
            prompt = task['prompt']
            schedule_type = task['schedule_type']
            
            print(f"🚀 Executing scheduled task [{task_id}]: {task['title']}")
            
            # اجرای پرامپت از طریق ایجنت هوش مصنوعی
            try:
                if self.agent_runner_callback:
                    self.agent_runner_callback(prompt)
            except Exception as ex:
                print(f"❌ Failed to run agent for task {task_id}: {ex}")
            
            # به‌روزرسانی زمان اجرای بعدی بر اساس نوع زمان‌بندی
            if schedule_type == 'once':
                cursor.execute("UPDATE tasks SET enabled = 0, last_run = ? WHERE id = ?", (now_str, task_id))
            elif schedule_type == 'daily':
                from datetime import timedelta
                next_time = (datetime.now() + timedelta(days=1)).strftime('%Y-%m-%d %H:%M:%S')
                cursor.execute("UPDATE tasks SET last_run = ?, next_run = ? WHERE id = ?", (now_str, next_time, task_id))
            elif schedule_type == 'interval':
                from datetime import timedelta
                minutes = int(task['repeat_interval'] or 60)
                next_time = (datetime.now() + timedelta(minutes=minutes)).strftime('%Y-%m-%d %H:%M:%S')
                cursor.execute("UPDATE tasks SET last_run = ?, next_run = ? WHERE id = ?", (now_str, next_time, task_id))
                
            conn.commit()
            
        conn.close()

    def stop(self):
        self._stop_event.set()


# ---------- TOOL SCHEMAS (چیزی که به مدل معرفی می‌شود) ----------

TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "خواندن محتوای کامل یک فایل از روی سیستم، همراه با تعداد خطوط و لیست توابع/کلاس‌های آن.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {"type": "string", "description": "مسیر کامل فایل، مثل C:\\AI_Assistant\\app\\main.py"}
                },
                "required": ["file_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "write_file",
            "description": "نوشتن یا جایگزینی محتوای یک فایل روی سیستم.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {"type": "string"},
                    "content": {"type": "string"}
                },
                "required": ["file_path", "content"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "web_search",
            "description": "جستجو در اینترنت برای اطلاعات به‌روز.",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string"}
                },
                "required": ["query"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "ssh_exec",
            "description": "اجرای یک دستور روی سرور راه دور از طریق SSH.",
            "parameters": {
                "type": "object",
                "properties": {
                    "hostname": {"type": "string"},
                    "username": {"type": "string"},
                    "password": {"type": "string"},
                    "command": {"type": "string"},
                    "port": {"type": "integer"}
                },
                "required": ["hostname", "username", "password", "command"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "open_browser",
            "description": "باز کردن مرورگر کروم واقعی روی سیستم و رفتن به یک آدرس وب مشخص (URL). یک پنجره مرورگر قابل‌مشاهده باز می‌شود. علاوه بر متن صفحه، یک اسکرین‌شات هم گرفته و توسط یک مدل تصویری (vision) تحلیل می‌شود تا ایجنت بفهمد صفحه از نظر ظاهری/بصری چه شکلی است (دکمه‌ها، فرم‌ها، چیدمان و غیره).",
            "parameters": {
                "type": "object",
                "properties": {
                    "url": {
                        "type": "string",
                        "description": "آدرس کامل وبسایت مورد نظر، مثل https://google.com"
                    },
                    "see_page": {
                        "type": "boolean",
                        "description": "اگر true باشد (پیش‌فرض)، از صفحه اسکرین‌شات گرفته و با مدل تصویری تحلیل می‌شود. اگر فقط متن صفحه کافی است، false بگذارید تا سریع‌تر اجرا شود."
                    }
                },
                "required": ["url"]
            }
        }
    }
]


# ---------- TOOL EXECUTION ----------

def execute_tool(name, args):
    try:
        if name == "read_file":
            path = args["file_path"]
            content = read_file(path)

            if content.startswith("خطا"):
                return content

            lines = content.splitlines()
            functions, classes = [], []
            for line in lines:
                s = line.strip()
                if s.startswith("def "):
                    functions.append(s.split("(")[0].replace("def ", ""))
                if s.startswith("class "):
                    classes.append(s.split(":")[0].replace("class ", ""))

            return (
                f"FILE_PATH: {path}\n"
                f"LINE_COUNT: {len(lines)}\n"
                f"FUNCTIONS: {', '.join(functions) if functions else 'ندارد'}\n"
                f"CLASSES: {', '.join(classes) if classes else 'ندارد'}\n\n"
                f"FILE_CONTENT:\n{content}"
            )
        
        elif name == "open_browser":
            return open_browser(args.get("url", "https://google.com"), see_page=args.get("see_page", True))

        elif name == "web_search":
            return web_search(args["query"])

        elif name == "write_file":
            return write_file(args["file_path"], args["content"])

        elif name == "ssh_exec":
            return ssh_exec(
                args["hostname"], args["username"], args["password"],
                args["command"], args.get("port", 22)
            )

        else:
            return f"ابزار ناشناخته: {name}"

    except Exception as e:
        return f"خطا در اجرای ابزار {name}: {e}"


# ---------- SYSTEM ----------

SYSTEM = """
تو یک دستیار هوشمند فارسی هستی که به ابزارها دسترسی داری.

قوانین مهم:
- اگر برای پاسخ به سوال کاربر نیاز به خواندن فایل، جستجوی وب، نوشتن فایل یا اجرای دستور SSH داری،
  حتماً از ابزار مربوطه استفاده کن. هرگز محتوای فایل یا نتیجه جستجو را حدس نزن.
- بعد از دریافت نتیجه‌ی یک ابزار، فقط بر اساس همان نتیجه‌ی واقعی پاسخ بده.
- اگر LINE_COUNT در نتیجه‌ی read_file وجود دارد، همان عدد دقیق را گزارش کن.
- وقتی اطلاعات کافی برای پاسخ نهایی داری، دیگر ابزار صدا نزن و مستقیم پاسخ نهایی فارسی بده.
"""


# ---------- OLLAMA CALL ----------

def call_ollama(messages, use_tools=True):
    payload = {
        "model": MODEL,
        "messages": messages,
        "stream": False,
        "options": {
            "temperature": 0.1,
            "num_ctx": 16384,
            "num_predict": 1024
        }
    }
    if use_tools:
        payload["tools"] = TOOLS

    r = requests.post(OLLAMA_URL, json=payload, timeout=600)
    r.raise_for_status()
    return r.json()["message"]


# ---------- AGENT LOOP ----------

def extract_attachment_text(file_path):
    ext = file_path.suffix.lower()

    try:
        if ext in (".txt", ".md", ".csv", ".json", ".py", ".log", ".html", ".xml", ".ini", ".yaml", ".yml"):
            for enc in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
                try:
                    return file_path.read_text(encoding=enc)
                except Exception:
                    continue
            return file_path.read_bytes().decode("utf-8", errors="replace")

        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            reader = PdfReader(str(file_path))
            return "\n".join((p.extract_text() or "") for p in reader.pages)

        if ext == ".docx":
            import docx
            doc = docx.Document(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs)

        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    parts.append(" | ".join("" if c is None else str(c) for c in row))
            return "\n".join(parts)

        if ext in (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"):
            return "[این یک فایل تصویری است. محتوای متنی قابل استخراج مستقیم ندارد.]"

        return file_path.read_bytes().decode("utf-8", errors="replace")

    except Exception as e:
        return f"خطا در استخراج محتوای فایل: {e}"


def run_agent(user_message, attachment=None):
    if attachment:
        try:
            from pathlib import Path
            file_path = Path("uploads") / Path(attachment).name
            if file_path.exists():
                file_text = extract_attachment_text(file_path)
                MAX_FILE_CHARS = 50000
                if len(file_text) > MAX_FILE_CHARS:
                    file_text = file_text[:MAX_FILE_CHARS] + "\n\n[محتوا طولانی بود و فقط بخش ابتدایی نمایش داده شد]"
                user_message = user_message + "\n\n========== محتوای فایل پیوست‌شده ==========\nنام فایل: " + file_path.name + "\n\n" + file_text + "\n========== پایان فایل =========="
            else:
                user_message += "\n\n[فایل پیوست‌شده پیدا نشد: " + str(attachment) + "]"
        except Exception as e:
            user_message += "\n\n[خطا در خواندن فایل پیوست‌شده: " + str(e) + "]"

    system_msg = {"role": "system", "content": SYSTEM + "\n\nحافظه دائمی:\n" + get_memory()}
    messages = [system_msg] + list(conversation_history) + [{"role": "user", "content": user_message}]
    steps_log = []

    for step in range(MAX_AGENT_STEPS):
        msg = call_ollama(messages)
        tool_calls = msg.get("tool_calls")

        if not tool_calls:
            final_answer = msg.get("content", "")
            conversation_history.append({"role": "user", "content": user_message})
            conversation_history.append({"role": "assistant", "content": final_answer})
            del conversation_history[:-MAX_HISTORY_MESSAGES]
            return final_answer, steps_log

        messages.append(msg)
        for call in tool_calls:
            fn_name = call["function"]["name"]
            fn_args = call["function"]["arguments"]

            if isinstance(fn_args, str):
                try:
                    fn_args = json.loads(fn_args)
                except Exception:
                    fn_args = {}

            result = execute_tool(fn_name, fn_args)
            is_error = isinstance(result, str) and result.startswith("خطا")

            steps_log.append({
                "tool": fn_name,
                "args": fn_args,
                "result_len": len(result) if result else 0,
                "is_error": is_error
            })

            tool_message = "⚠️ خطا در اجرای ابزار: " + result if is_error else result
            messages.append({"role": "tool", "content": tool_message})

    return "متاسفانه پس از چند مرحله نتوانستم به پاسخ نهایی برسم.", steps_log


# ---------- HTML & WORKSPACE FRONTEND ----------

HTML = """
<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>دستیار هوش مصنوعی - AI Workspace</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Vazirmatn:wght@300;400;500;600;700&display=swap');

*{ box-sizing:border-box; margin:0; padding:0; }

:root {
  --bg-main: #edf5f0;
  --panel-bg: rgba(255, 255, 255, 0.75);
  --panel-border: rgba(255, 255, 255, 0.85);
  --text-main: #26352d;
  --text-muted: #718078;
  --primary: #54846c;
  --primary-hover: #4b7c63;
  --primary-light: #e8f5ed;
  --accent: #72a98b;
  --shadow: 0 12px 35px rgba(45, 75, 60, 0.08);
}

body{
  min-height:100vh;
  font-family:'Vazirmatn',Tahoma,Arial,sans-serif;
  background:var(--bg-main);
  background-image: linear-gradient(rgba(247,249,245,.12),rgba(247,249,245,.12)), url('/static/background.jpg');
  background-size:cover;
  background-position:center;
  background-attachment:fixed;
  color:var(--text-main);
  display: flex;
}

/* Layout App Workspace */
.app-container {
  display: flex;
  width: 100vw;
  height: 100vh;
  overflow: hidden;
}

/* Sidebar */
aside {
  width: 280px;
  background: rgba(255, 255, 255, 0.85);
  backdrop-filter: blur(20px);
  -webkit-backdrop-filter: blur(20px);
  border-left: 1px solid rgba(255, 255, 255, 0.9);
  display: flex;
  flex-direction: column;
  box-shadow: -4px 0 25px rgba(45, 75, 60, 0.05);
  z-index: 10;
}

.sidebar-header {
  padding: 22px 20px;
  display: flex;
  align-items: center;
  gap: 14px;
  border-bottom: 1px solid rgba(0,0,0,0.05);
}

.logo{
  width:44px;
  height:44px;
  border-radius:14px;
  display:flex;
  align-items:center;
  justify-content:center;
  font-size:22px;
  font-weight:700;
  color:white;
  background:linear-gradient(145deg,#72a98b,#4f8069);
  box-shadow:0 6px 15px rgba(69,123,94,.28);
}

.brand-title{
  font-size:17px;
  font-weight:700;
  color:#315846;
}

.brand-subtitle{
  font-size:11px;
  color:var(--text-muted);
}

.sidebar-menu {
  padding: 15px 12px;
  display: flex;
  flex-direction: column;
  gap: 6px;
  overflow-y: auto;
  flex: 1;
}

.menu-item {
  display: flex;
  align-items: center;
  gap: 12px;
  padding: 12px 16px;
  border-radius: 14px;
  font-size: 14px;
  font-weight: 500;
  color: #3e4e45;
  cursor: pointer;
  transition: all 0.2s ease;
  border: 1px solid transparent;
}

.menu-item:hover {
  background: rgba(114, 169, 139, 0.1);
  color: #2e5440;
}

.menu-item.active {
  background: linear-gradient(135deg, #e8f5ed, #dcefe5);
  color: #315846;
  border-color: rgba(114, 169, 139, 0.25);
  font-weight: 600;
  box-shadow: 0 4px 12px rgba(84, 132, 108, 0.08);
}

/* Main Content Area */
main {
  flex: 1;
  display: flex;
  flex-direction: column;
  height: 100vh;
  overflow: hidden;
}

header{
  display:flex;
  align-items:center;
  justify-content:space-between;
  padding:18px 32px;
  background:var(--panel-bg);
  backdrop-filter:blur(18px);
  -webkit-backdrop-filter:blur(18px);
  border-bottom:1px solid var(--panel-border);
  box-shadow:0 4px 20px rgba(45,75,60,.04);
}

.header-title {
  font-size: 18px;
  font-weight: 700;
  color: #315846;
}

.header-badge {
  font-size: 12px;
  padding: 6px 14px;
  background: var(--primary-light);
  border: 1px solid #b9d9c5;
  color: #315f46;
  border-radius: 20px;
  font-weight: 500;
}

/* Views Container */
.view-panel {
  display: none;
  flex: 1;
  overflow-y: auto;
  padding: 28px 32px;
}

.view-panel.active {
  display: flex;
  flex-direction: column;
}

/* Dashboard View */
.welcome-banner {
  background: linear-gradient(135deg, rgba(255,255,255,0.9), rgba(232,245,237,0.8));
  border: 1px solid rgba(255,255,255,0.9);
  padding: 28px 32px;
  border-radius: 24px;
  margin-bottom: 24px;
  box-shadow: var(--shadow);
}

.welcome-banner h1 {
  font-size: 24px;
  font-weight: 700;
  color: #315846;
  margin-bottom: 8px;
}

.welcome-banner p {
  color: var(--text-muted);
  font-size: 14px;
}

.quick-cards {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
  gap: 16px;
  margin-bottom: 28px;
}

.quick-card {
  background: rgba(255, 255, 255, 0.7);
  border: 1px solid rgba(255, 255, 255, 0.9);
  padding: 20px;
  border-radius: 18px;
  cursor: pointer;
  transition: all 0.2s;
  box-shadow: 0 4px 15px rgba(0,0,0,0.02);
  display: flex;
  align-items: center;
  gap: 14px;
}

.quick-card:hover {
  transform: translateY(-3px);
  background: rgba(255, 255, 255, 0.95);
  box-shadow: 0 8px 25px rgba(84, 132, 108, 0.12);
}

.quick-card-icon {
  width: 42px;
  height: 42px;
  border-radius: 12px;
  background: var(--primary-light);
  color: var(--primary);
  display: flex;
  align-items: center;
  justify-content: center;
  font-size: 18px;
  font-weight: bold;
}

.quick-card-info h3 {
  font-size: 14px;
  font-weight: 600;
  color: #33423a;
  margin-bottom: 2px;
}

.quick-card-info span {
  font-size: 12px;
  color: var(--text-muted);
}

/* Sections & Cards */
.section-grid {
  display: grid;
  grid-template-columns: 2fr 1fr;
  gap: 20px;
}

@media(max-width: 900px) {
  .section-grid { grid-template-columns: 1fr; }
}

.dashboard-card {
  background: rgba(255, 255, 255, 0.75);
  border: 1px solid rgba(255, 255, 255, 0.9);
  border-radius: 20px;
  padding: 22px;
  box-shadow: var(--shadow);
  margin-bottom: 20px;
}

.dashboard-card h3 {
  font-size: 16px;
  font-weight: 600;
  color: #315846;
  margin-bottom: 16px;
  display: flex;
  align-items: center;
  justify-content: space-between;
}

.activity-item, .project-mini-card, .task-mini-item {
  padding: 12px 14px;
  border-radius: 12px;
  background: rgba(247, 249, 245, 0.6);
  margin-bottom: 10px;
  border: 1px solid rgba(0,0,0,0.03);
  font-size: 13px;
  display: flex;
  justify-content: space-between;
  align-items: center;
}

/* Chat View Styling */
#chat-container {
  display: flex;
  flex-direction: column;
  flex: 1;
  overflow: hidden;
}

#chat{
  flex: 1;
  overflow-y: auto;
  padding: 20px;
  background: rgba(255,255,255,.62);
  backdrop-filter:blur(18px);
  -webkit-backdrop-filter:blur(18px);
  border:1px solid rgba(255,255,255,.80);
  border-radius:22px;
  box-shadow: inset 0 1px 0 rgba(255,255,255,.9);
  margin-bottom: 16px;
}

#chat:empty:before{
  content:'سلام!\A چطور می‌تونم کمکتون کنم؟\A پیامتون رو بنویسید و بفرستید';
  white-space:pre;
  display:block;
  text-align:center;
  margin-top:130px;
  line-height:2.2;
  font-size:16px;
  color:#7a8981;
}

.msg{
  max-width:82%;
  padding:12px 16px;
  margin:10px 0;
  border-radius:16px;
  white-space:pre-wrap;
  line-height:1.9;
  font-size:14px;
  box-shadow:0 4px 12px rgba(50,70,60,.04);
}

.user{
  margin-right:0;
  margin-left:auto;
  background:linear-gradient(135deg,#dcefe5,#e8f4ed);
  color:#315646;
  border-bottom-right-radius:4px;
  border:1px solid rgba(105,153,125,.15);
}

.ai{
  margin-left:0;
  margin-right:auto;
  background:rgba(255,251,243,.91);
  color:#3e4943;
  border-bottom-left-radius:4px;
  border:1px solid rgba(181,161,123,.14);
}

#inputbox{
  display:flex;
  gap:10px;
  align-items:stretch;
  direction:ltr;
  background: rgba(255,255,255,0.7);
  padding: 12px;
  border-radius: 20px;
  border: 1px solid rgba(255,255,255,0.9);
}

textarea{
  flex:1;
  resize:none;
  border:1px solid rgba(125,153,137,.28);
  border-radius:16px;
  padding:12px 16px;
  font-family:'Vazirmatn',Tahoma,Arial,sans-serif;
  font-size:14px;
  line-height:1.7;
  background:rgba(255,255,255,.85);
  color:#33423a;
  outline:none;
  direction:rtl;
}

textarea:focus{
  border-color:#78a88e;
  box-shadow:0 0 0 3px rgba(112,160,133,.12);
}

button{
  border:0;
  border-radius:14px;
  background:linear-gradient(145deg,#6fa489,#54846c);
  color:white;
  font-family:'Vazirmatn',Tahoma,Arial,sans-serif;
  font-size:14px;
  font-weight:500;
  cursor:pointer;
  padding: 0 22px;
  box-shadow:0 6px 15px rgba(67,119,91,.22);
  transition:.2s;
}

button:hover{
  background:linear-gradient(145deg,#639b7e,#4b7c63);
  transform:translateY(-1px);
}

#stopBtn.active{
  cursor:pointer !important;
  opacity:1 !important;
  background:linear-gradient(145deg,#e57373,#c0392b) !important;
  color:white !important;
}

/* Projects & Plugins Grid */
.grid-cards {
  display: grid;
  grid-template-columns: repeat(auto-fill, minmax(260px, 1fr));
  gap: 20px;
}

.card-item {
  background: rgba(255, 255, 255, 0.8);
  border: 1px solid rgba(255, 255, 255, 0.95);
  border-radius: 20px;
  padding: 20px;
  box-shadow: var(--shadow);
  display: flex;
  flex-direction: column;
  justify-content: space-between;
  transition: transform 0.2s;
}

.card-item:hover {
  transform: translateY(-2px);
}

.card-item h4 {
  font-size: 16px;
  font-weight: 600;
  color: #315846;
  margin-bottom: 8px;
}

.card-item p {
  font-size: 13px;
  color: var(--text-muted);
  line-height: 1.7;
  margin-bottom: 16px;
}

/* Modal Styling for Tasks */
.modal {
  display: none; 
  position: fixed; 
  z-index: 100; 
  left: 0; top: 0; width: 100%; height: 100%; 
  background-color: rgba(0,0,0,0.4); 
  backdrop-filter: blur(4px);
  align-items: center; justify-content: center;
}
.modal-content {
  background: #fff; 
  padding: 24px; 
  border-radius: 20px; 
  width: 420px; 
  box-shadow: var(--shadow);
  display: flex; flex-direction: column; gap: 14px;
}
.modal-content h3 { color: #315846; font-size: 16px; margin-bottom: 4px; }
.modal-content input, .modal-content select, .modal-content textarea {
  width: 100%; padding: 10px 14px; border-radius: 10px; border: 1px solid #b9d9c5;
  font-family: 'Vazirmatn'; font-size: 13px; outline: none; background: #fafafa;
}

/* Settings Form */
.settings-group {
  background: rgba(255, 255, 255, 0.75);
  border: 1px solid rgba(255, 255, 255, 0.9);
  border-radius: 20px;
  padding: 24px;
  margin-bottom: 20px;
  box-shadow: var(--shadow);
}

.settings-group h3 {
  font-size: 16px;
  font-weight: 600;
  color: #315846;
  margin-bottom: 16px;
  border-bottom: 1px solid rgba(0,0,0,0.05);
  padding-bottom: 10px;
}

.form-row {
  display: flex;
  align-items: center;
  justify-content: space-between;
  margin-bottom: 14px;
  font-size: 14px;
}

.form-row label {
  color: #3e4e45;
  font-weight: 500;
}

.form-row input[type="text"], .form-row select {
  padding: 8px 14px;
  border-radius: 10px;
  border: 1px solid #b9d9c5;
  background: #fff;
  font-family: 'Vazirmatn';
  font-size: 13px;
  width: 220px;
  outline: none;
}

.form-row input[type="checkbox"] {
  width: 18px;
  height: 18px;
  accent-color: var(--primary);
}

/* Switch toggle */
.switch {
  position: relative;
  display: inline-block;
  width: 48px;
  height: 24px;
}
.switch input { opacity: 0; width: 0; height: 0; }
.slider {
  position: absolute; cursor: pointer; top: 0; left: 0; right: 0; bottom: 0;
  background-color: #ccc; transition: .3s; border-radius: 24px;
}
.slider:before {
  position: absolute; content: ""; height: 18px; width: 18px; left: 3px; bottom: 3px;
  background-color: white; transition: .3s; border-radius: 50%;
}
input:checked + .slider { background-color: var(--primary); }
input:checked + .slider:before { transform: translateX(24px); }

/* Responsive design */
@media(max-width: 768px) {
  aside { width: 70px; }
  .brand-title, .brand-subtitle, .menu-item span { display: none; }
  .sidebar-header { justify-content: center; padding: 15px; }
  .menu-item { justify-content: center; padding: 12px; }
}
</style>
</head>
<body>

<div class="app-container">
  <!-- Sidebar -->
  <aside>
    <div class="sidebar-header">
      <div class="logo">AI</div>
      <div>
        <div class="brand-title">دستیار هوشمند</div>
        <div class="brand-subtitle">Workspace v2.5</div>
      </div>
    </div>
    <div class="sidebar-menu">
      <div class="menu-item active" onclick="switchView('dashboard', this)">
        <span>📊</span> داشبورد اصلی
      </div>
      <div class="menu-item" onclick="switchView('chat', this)">
        <span>💬</span> گفتگوها و چت
      </div>
      <div class="menu-item" onclick="switchView('projects', this)">
        <span>📁</span> پروژه‌ها
      </div>
      <div class="menu-item" onclick="switchView('plugins', this)">
        <span>🧩</span> پلاگین‌ها و ابزارها
      </div>
      <div class="menu-item" onclick="switchView('tasks', this); loadTasks();">
        <span>⏰</span> وظایف برنامه‌ریزی‌شده
      </div>
      <div class="menu-item" onclick="switchView('settings', this)">
        <span>⚙️</span> تنظیمات سیستم
      </div>
    </div>
  </aside>

  <!-- Main Content Area -->
  <main>
    <!-- Header -->
    <header>
      <div id="headerTitle" class="header-title">داشبورد اصلی</div>
      <div class="header-badge">Qwen 2.5 7B • فعال</div>
    </header>

    <!-- 1. Dashboard View -->
    <div id="view-dashboard" class="view-panel active">
      <div class="welcome-banner">
        <h1>سلام، محمدرضا 👋</h1>
        <p>امروز چه کاری می‌توانم برایت انجام دهم؟ از بخش‌های زیر برای مدیریت هوشمند کارهایت استفاده کن.</p>
      </div>

      <div class="quick-cards">
        <div class="quick-card" onclick="switchView('chat', document.querySelectorAll('.menu-item')[1])">
          <div class="quick-card-icon">💬</div>
          <div class="quick-card-info">
            <h3>گفتگوی جدید</h3>
            <span>شروع گفتگو با هوش مصنوعی</span>
          </div>
        </div>
        <div class="quick-card" onclick="switchView('projects', document.querySelectorAll('.menu-item')[2])">
          <div class="quick-card-icon">📁</div>
          <div class="quick-card-info">
            <h3>پروژه جدید</h3>
            <span>مدیریت فایل‌ها و پروژه‌ها</span>
          </div>
        </div>
        <div class="quick-card" onclick="switchView('plugins', document.querySelectorAll('.menu-item')[3])">
          <div class="quick-card-icon">🔍</div>
          <div class="quick-card-info">
            <h3>جستجوی وب</h3>
            <span>دستیابی به اطلاعات آنلاین</span>
          </div>
        </div>
        <div class="quick-card" onclick="switchView('tasks', document.querySelectorAll('.menu-item')[4]); loadTasks();">
          <div class="quick-card-icon">⏰</div>
          <div class="quick-card-info">
            <h3>وظیفه جدید</h3>
            <span>زمان‌بندی کارهای خودکار</span>
          </div>
        </div>
      </div>

      <div class="section-grid">
        <div>
          <div class="dashboard-card">
            <h3>📁 پروژه‌های اخیر</h3>
            <div class="project-mini-card">
              <span>سیستم مدیریت بالینی پزشکی</span>
              <span style="color:var(--text-muted)">آخرین فعالیت: امروز</span>
            </div>
            <div class="project-mini-card">
              <span>مرکز مشاوره آرزو - پروتکل درمانی</span>
              <span style="color:var(--text-muted)">آخرین فعالیت: دیروز</span>
            </div>
          </div>

          <div class="dashboard-card">
            <h3>📈 فعالیت‌های اخیر</h3>
            <div class="activity-item">
              <span>اجرای ابزار خواندن فایل (read_file)</span>
              <span style="color:var(--text-muted)">۱ ساعت پیش</span>
            </div>
            <div class="activity-item">
              <span>بروزرسانی حافظه دائمی دستیار</span>
              <span style="color:var(--text-muted)">۳ ساعت پیش</span>
            </div>
          </div>
        </div>

        <div>
          <div class="dashboard-card">
            <h3>⏰ وظایف بعدی</h3>
            <div class="task-mini-item">
              <span>گزارش ماهانه فعالیت بالینی</span>
              <span style="color:var(--primary); font-weight:600;">سیستم خودکار</span>
            </div>
            <div class="task-mini-item">
              <span>بررسی تسک‌های پس‌زمینه</span>
              <span style="color:var(--primary); font-weight:600;">فعال</span>
            </div>
          </div>

          <div class="dashboard-card">
            <h3>⚙️ وضعیت سیستم</h3>
            <div style="font-size: 13px; color: var(--text-muted); line-height: 1.8;">
              <div>مدل فعال: <b style="color:var(--text-main)">Qwen 2.5 7B</b></div>
              <div>وضعیت سرور: <b style="color:#2e7d32">متصل و پایدار</b></div>
              <div>حافظه دائمی: <b style="color:var(--text-main)">فعال (SQLite)</b></div>
            </div>
          </div>
        </div>
      </div>
    </div>

    <!-- 2. Chat View -->
    <div id="view-chat" class="view-panel">
      <div id="chat-container">
        <div id="chat"></div>
        <div id="inputbox">
          <div style="display:flex;align-items:center;gap:10px;">
            <label for="fileInput" style="cursor:pointer;padding:9px 14px;border-radius:12px;background:#e8f5ed;border:1px solid #b9d9c5;color:#315f46;font-size:13px;font-weight:600;">
              📎 الصاق فایل
            </label>
            <input id="fileInput" type="file" accept=".pdf,.docx,.txt,.csv,.json,.xlsx,.xls,.md,.py,.log,.html,.xml,.jpg,.jpeg,.png" style="display:none;">
            <span id="fileStatus" style="font-size:12px;color:var(--text-muted)">فایلی انتخاب نشده</span>
          </div>

          <textarea id="input" rows="2" placeholder="پیام خود را بنویسید..."></textarea>

          <div style="display:flex;gap:8px;">
            <button onclick="send()" id="sendBtn">ارسال</button>
            <button onclick="stopGeneration()" id="stopBtn" title="توقف فعالیت" disabled style="background:#f3e9e9;color:#b94a4a;opacity:.55;">توقف</button>
          </div>
        </div>
      </div>
    </div>

    <!-- 3. Projects View -->
    <div id="view-projects" class="view-panel">
      <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:20px;">
        <h2 style="font-size:18px; color:#315846;">مدیریت پروژه‌ها</h2>
        <button onclick="alert('قابلیت ایجاد پروژه جدید آماده اتصال به دیتابیس است.')">+ پروژه جدید</button>
      </div>
      <div class="grid-cards">
        <div class="card-item">
          <div>
            <h4>سیستم مدیریت پزشکی دانشگاه</h4>
            <p>توسعه اسکریپت‌های پایتون و Flask برای مدیریت شیفت‌ها، ساعات کلینیک و گزارش‌گیری.</p>
          </div>
          <span style="font-size:12px; color:var(--primary); font-weight:600;">آخرین فعالیت: امروز</span>
        </div>
        <div class="card-item">
          <div>
            <h4>مرکز مشاوره آرزو</h4>
            <p>طراحی و تدوین پروتکل‌های درمانی و مدیریت اطلاعات مراجعان و ابزارهای سنجش.</p>
          </div>
          <span style="font-size:12px; color:var(--primary); font-weight:600;">آخرین فعالیت: دیروز</span>
        </div>
      </div>
    </div>

    <!-- 4. Plugins View -->
    <div id="view-plugins" class="view-panel">
      <h2 style="font-size:18px; color:#315846; margin-bottom:20px;">پلاگین‌ها و ابزارهای دستیار</h2>
      <div class="grid-cards">
        <div class="card-item">
          <div>
            <h4>جستجوی وب</h4>
            <p>دستیابی به اطلاعات آنلاین و به‌روز اینترنتی از طریق موتور جستجو.</p>
          </div>
          <div style="display:flex; justify-content:space-between; align-items:center;">
            <span style="font-size:12px; color:#2e7d32; font-weight:600;">فعال</span>
            <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
          </div>
        </div>
        <div class="card-item">
          <div>
            <h4>تحلیل فایل و اسناد</h4>
            <p>خواندن، استخراج متن و پردازش فایل‌های PDF، Word، Excel و متنی.</p>
          </div>
          <div style="display:flex; justify-content:space-between; align-items:center;">
            <span style="font-size:12px; color:#2e7d32; font-weight:600;">فعال</span>
            <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
          </div>
        </div>
        <div class="card-item">
          <div>
            <h4>حافظه دائمی (Memory)</h4>
            <p>ذخیره و بازیابی اطلاعات و نکات مهم در پایگاه داده SQLite.</p>
          </div>
          <div style="display:flex; justify-content:space-between; align-items:center;">
            <span style="font-size:12px; color:#2e7d32; font-weight:600;">فعال</span>
            <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
          </div>
        </div>
        <div class="card-item">
          <div>
            <h4>اجرای دستورات SSH</h4>
            <p>اتصال و اجرای دستورات روی سرورهای راه دور به صورت امن.</p>
          </div>
          <div style="display:flex; justify-content:space-between; align-items:center;">
            <span style="font-size:12px; color:#2e7d32; font-weight:600;">فعال</span>
            <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
          </div>
        </div>
      </div>
    </div>

    <!-- 5. Scheduled Tasks View -->
    <div id="view-tasks" class="view-panel">
      <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:20px;">
        <h2 style="font-size:18px; color:#315846;">وظایف برنامه‌ریزی‌شده</h2>
        <button onclick="openTaskModal()">+ وظیفه جدید</button>
      </div>
      <div class="dashboard-card" id="tasksListContainer">
        <p style="color: var(--text-muted); font-size: 13px;">در حال بارگذاری وظایف...</p>
      </div>
    </div>

    <!-- 6. Settings View -->
    <div id="view-settings" class="view-panel">
      <h2 style="font-size:18px; color:#315846; margin-bottom:20px;">تنظیمات سیستم</h2>

      <div class="settings-group">
        <h3>شخصی‌سازی و هوش مصنوعی</h3>
        <div class="form-row">
          <label>نام کاربر</label>
          <input type="text" id="setUserName" value="محمدرضا">
        </div>
        <div class="form-row">
          <label>مدل هوش مصنوعی</label>
          <input type="text" id="setModelName" value="qwen2.5:7b" disabled>
        </div>
        <div class="form-row">
          <label>دمای پاسخ (Temperature)</label>
          <input type="text" id="setTemp" value="0.1">
        </div>
      </div>

      <div class="settings-group">
        <h3>حریم خصوصی و حافظه</h3>
        <div class="form-row">
          <label>ذخیره تاریخچه مکالمات</label>
          <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
        </div>
        <div class="form-row">
          <label>حافظه دائمی فعال باشد</label>
          <label class="switch"><input type="checkbox" checked><span class="slider"></span></label>
        </div>
      </div>

      <div style="display:flex; justify-content:flex-end;">
        <button onclick="saveSettings()">ذخیره تنظیمات</button>
      </div>
    </div>
  </main>
</div>

<!-- Modal برای ساخت تسک جدید -->
<div id="taskModal" class="modal">
  <div class="modal-content">
    <h3>تعریف وظیفه زمان‌بندی‌شده جدید</h3>
    <input type="text" id="taskTitle" placeholder="عنوان وظیفه (مثلا: گزارش وضعیت)">
    <textarea id="taskPrompt" rows="3" placeholder="پرامپتی که باید خودکار اجرا شود..."></textarea>
    <select id="taskScheduleType" onchange="toggleScheduleFields()">
      <option value="once">یک‌بار در زمان مشخص</option>
      <option value="daily">روزانه (هر ۲۴ ساعت)</option>
      <option value="interval">دوره‌ای (بر حسب دقیقه)</option>
    </select>
    <div id="runAtField">
      <label style="font-size:12px; color:var(--text-muted)">زمان اجرا (YYYY-MM-DD HH:MM:SS)</label>
      <input type="text" id="taskRunAt" placeholder="2026-06-01 08:00:00">
    </div>
    <div id="intervalField" style="display:none;">
      <label style="font-size:12px; color:var(--text-muted)">فاصله زمانی (دقیقه)</label>
      <input type="number" id="taskInterval" value="60">
    </div>
    <div style="display:flex; justify-content:flex-end; gap:8px; margin-top:10px;">
      <button onclick="closeTaskModal()" style="background:#ccc;">انصراف</button>
      <button onclick="createTask()">ذخیره و فعال‌سازی</button>
    </div>
  </div>
</div>

<script>
// SPA Navigation logic
function switchView(viewName, element) {
  document.querySelectorAll('.view-panel').forEach(panel => panel.classList.remove('active'));
  document.querySelectorAll('.menu-item').forEach(item => item.classList.remove('active'));

  document.getElementById('view-' + viewName).classList.add('active');
  if(element) element.classList.add('active');

  const titles = {
    'dashboard': 'داشبورد اصلی',
    'chat': 'گفتگو با دستیار هوشمند',
    'projects': 'مدیریت پروژه‌ها',
    'plugins': 'پلاگین‌ها و ابزارها',
    'tasks': 'وظایف برنامه‌ریزی‌شده',
    'settings': 'تنظیمات سیستم'
  };
  document.getElementById('headerTitle').textContent = titles[viewName] || 'دستیار هوشمند';
}

// Chat Functionality
function add(text,type){
  const d=document.createElement("div");
  d.className="msg "+type;
  d.textContent=text;
  const chatBox = document.getElementById("chat");
  chatBox.appendChild(d);
  chatBox.scrollTop = chatBox.scrollHeight;
}

let selectedFile = "";

document.getElementById("fileInput").addEventListener("change", async function () {
    const file = this.files[0];
    const status = document.getElementById("fileStatus");

    if (!file) {
        selectedFile = "";
        status.textContent = "فایلی انتخاب نشده";
        return;
    }

    status.textContent = "در حال آپلود...";
    const formData = new FormData();
    formData.append("file", file);

    try {
        const response = await fetch("/upload", { method: "POST", body: formData });
        const data = await response.json();

        if (data.success) {
            selectedFile = data.filename;
            status.textContent = "✅ " + data.filename;
            add("فایل پیوست شد: " + data.filename, "user");
        } else {
            selectedFile = "";
            status.textContent = "خطا";
            add("خطا در آپلود: " + (data.error || "دلیل نامشخص"), "ai");
        }
    } catch (error) {
        selectedFile = "";
        status.textContent = "خطا در اتصال";
        add("خطا در ارتباط با سرور آپلود فایل", "ai");
    }
});

let currentController = null;

function setBusy(isBusy){
  const stopBtn = document.getElementById("stopBtn");
  const sendBtn = document.getElementById("sendBtn");
  if(isBusy){
    stopBtn.disabled = false;
    stopBtn.classList.add("active");
    sendBtn.disabled = true;
    sendBtn.style.opacity = ".6";
  } else {
    stopBtn.disabled = true;
    stopBtn.classList.remove("active");
    sendBtn.disabled = false;
    sendBtn.style.opacity = "1";
  }
}

function stopGeneration(){
  if(currentController){
    currentController.abort();
    currentController = null;
  }
  setBusy(false);
  const chat=document.getElementById("chat");
  if(chat.lastChild && chat.lastChild.textContent === "در حال پردازش..."){
    chat.removeChild(chat.lastChild);
  }
  add("فعالیت توسط کاربر متوقف شد.","ai");
}

async function send(){
  const box=document.getElementById("input");
  const text=box.value.trim();
  if(!text && !selectedFile) return;
  
  add(text || "[ارسال فایل پیوست]", "user");
  box.value="";
  add("در حال پردازش...","ai");
  setBusy(true);

  currentController = new AbortController();

  try{
    const r = await fetch("/chat",{
      method:"POST",
      headers:{"Content-Type":"application/json"},
      body:JSON.stringify({ message: text, attachment: selectedFile }),
      signal: currentController.signal
    });
    const data = await r.json();
    const chat = document.getElementById("chat");
    if(chat.lastChild && chat.lastChild.textContent === "در حال پردازش..."){
      chat.removeChild(chat.lastChild);
    }
    add(data.response || data.error,"ai");
  }catch(e){
    if(e.name !== "AbortError"){
      add("خطا: "+e,"ai");
    }
  }finally{
    setBusy(false);
    currentController = null;
    selectedFile = "";
    document.getElementById("fileStatus").textContent = "فایلی انتخاب نشده";
  }
}

document.getElementById("input").addEventListener("keydown",function(e){
  if(e.key==="Enter" && !e.shiftKey){ e.preventDefault(); send(); }
});

// Tasks Frontend Logic
async function loadTasks() {
    const container = document.getElementById("tasksListContainer");
    container.innerHTML = '<p style="color: var(--text-muted); font-size: 13px;">در حال بارگذاری...</p>';
    try {
        const res = await fetch('/api/tasks');
        const data = await res.json();
        if(data.status === 'success' && data.tasks.length > 0) {
            container.innerHTML = '';
            data.tasks.forEach(task => {
                const item = document.createElement('div');
                item.className = 'activity-item';
                item.innerHTML = `
                  <div>
                    <strong>${task.title}</strong>
                    <div style="font-size:12px; color:var(--text-muted)">نوع: ${task.schedule_type} | بعدی: ${task.next_run || 'تعیین نشده'}</div>
                  </div>
                  <div style="display:flex; gap:8px; align-items:center;">
                    <button onclick="runTaskManual(${task.id})" style="padding:4px 10px; font-size:12px; background:var(--primary);">اجرای فوری</button>
                    <label class="switch"><input type="checkbox" ${task.enabled ? 'checked' : ''} onchange="toggleTask(${task.id})"><span class="slider"></span></label>
                  </div>
                `;
                container.appendChild(item);
            });
        } else {
            container.innerHTML = '<p style="color: var(--text-muted); font-size: 13px;">هیچ وظیفه‌ای تعریف نشده است.</p>';
        }
    } catch(err) {
        container.innerHTML = '<p style="color: #c0392b; font-size: 13px;">خطا در بارگذاری وظایف.</p>';
    }
}

function openTaskModal() {
    document.getElementById("taskModal").style.display = "flex";
    const now = new Date();
    now.setMinutes(now.getMinutes() + 5);
    document.getElementById("taskRunAt").value = now.toISOString().slice(0, 19).replace('T', ' ');
}

function closeTaskModal() {
    document.getElementById("taskModal").style.display = "none";
}

function toggleScheduleFields() {
    const type = document.getElementById("taskScheduleType").value;
    document.getElementById("runAtField").style.display = type === 'interval' ? 'none' : 'block';
    document.getElementById("intervalField").style.display = type === 'interval' ? 'block' : 'none';
}

async function createTask() {
    const title = document.getElementById("taskTitle").value.trim();
    const prompt = document.getElementById("taskPrompt").value.trim();
    const schedule_type = document.getElementById("taskScheduleType").value;
    const run_at = document.getElementById("taskRunAt").value.trim();
    const repeat_interval = document.getElementById("taskInterval").value;

    if(!title || !prompt) {
        alert("لطفاً عنوان و پرامپت را وارد کنید.");
        return;
    }

    try {
        const res = await fetch('/api/tasks', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ title, prompt, schedule_type, run_at, repeat_interval })
        });
        const data = await res.json();
        if(data.status === 'success') {
            closeTaskModal();
            loadTasks();
            alert("وظیفه با موفقیت ثبت شد.");
        } else {
            alert("خطا در ثبت وظیفه");
        }
    } catch(e) {
        alert("خطا در ارتباط با سرور");
    }
}

async function toggleTask(id) {
    await fetch(`/api/tasks/${id}/toggle`, { method: 'POST' });
    loadTasks();
}

async function runTaskManual(id) {
    if(!confirm("آیا می‌خواهید این وظیفه هم‌اکنون اجرا شود؟")) return;
    try {
        const res = await fetch(`/api/tasks/${id}/run`, { method: 'POST' });
        const data = await res.json();
        if(data.status === 'success') {
            alert("وظیفه با موفقیت اجرا شد.");
            loadTasks();
        } else {
            alert("خطا: " + data.message);
        }
    } catch(e) {
        alert("خطا در اجرای دستی تسک");
    }
}

function saveSettings() {
  const userName = document.getElementById("setUserName").value;
  localStorage.setItem("assistant_username", userName);
  alert("تنظیمات با موفقیت ذخیره شد.");
}

window.addEventListener("DOMContentLoaded", () => {
  const savedName = localStorage.getItem("assistant_username");
  if(savedName) {
    document.getElementById("setUserName").value = savedName;
  }
});
</script>
</body>
</html>
"""


# ---------- ROUTES ----------

@app.route("/")
def home():
    return render_template_string(HTML)


@app.route("/chat", methods=["POST"])
def chat():
    try:
        user = request.json.get("message", "").strip()
        attachment = request.json.get("attachment", "").strip()
        if not user and not attachment:
            return jsonify({"response": "لطفاً پیام خود را وارد کنید."})

        if user.startswith("یادداشت کن ") or user.startswith("ذخیره کن "):
            prefix_len = len("یادداشت کن ") if user.startswith("یادداشت کن ") else len("ذخیره کن ")
            save_memory(user[prefix_len:].strip())
            return jsonify({"response": "در حافظه دائمی ذخیره شد."})

        if user == "حافظه":
            return jsonify({"response": get_memory()})

        if user in ("پاک کن تاریخچه", "شروع تازه"):
            conversation_history.clear()
            return jsonify({"response": "تاریخچه مکالمه پاک شد. گفتگوی جدید شروع شد."})

        answer, steps_log = run_agent(
            user,
            attachment if attachment else None
        )
        return jsonify({"response": answer, "steps": steps_log})

    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ---------- TASKS API ENDPOINTS ----------

@app.route('/api/tasks', methods=['GET'])
def get_tasks():
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM tasks ORDER BY id DESC")
        rows = cursor.fetchall()
        tasks = [dict(row) for row in rows]
        conn.close()
        return jsonify({"status": "success", "tasks": tasks})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/tasks', methods=['POST'])
def create_task():
    try:
        data = request.json
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        schedule_type = data.get('schedule_type', 'once')
        run_at = data.get('run_at')
        repeat_interval = data.get('repeat_interval')
        
        cursor.execute('''
            INSERT INTO tasks (title, description, prompt, schedule_type, run_at, repeat_interval, next_run)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (
            data.get('title'),
            data.get('description', ''),
            data.get('prompt'),
            schedule_type,
            run_at,
            repeat_interval,
            run_at if schedule_type != 'interval' else datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ))
        conn.commit()
        conn.close()
        return jsonify({"status": "success", "message": "Task created successfully."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/tasks/<int:task_id>/toggle', methods=['POST'])
def toggle_task(task_id):
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT enabled FROM tasks WHERE id = ?", (task_id,))
        row = cursor.fetchone()
        if row:
            new_status = 0 if row[0] == 1 else 1
            cursor.execute("UPDATE tasks SET enabled = ? WHERE id = ?", (new_status, task_id))
            conn.commit()
        conn.close()
        return jsonify({"status": "success"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/tasks/<int:task_id>/run', methods=['POST'])
def run_task_manually(task_id):
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("SELECT prompt FROM tasks WHERE id = ?", (task_id,))
        row = cursor.fetchone()
        conn.close()
        
        if row and task_scheduler_instance:
            task_scheduler_instance.agent_runner_callback(row['prompt'])
            return jsonify({"status": "success", "message": "Task executed manually."})
        return jsonify({"status": "error", "message": "Task not found or scheduler not active."}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# تعریف متغیر سراسری برای زمان‌بندی تسک‌ها
task_scheduler_instance = None


if __name__ == "__main__":
    # راه‌اندازی نخ پس‌زمینه زمان‌بندی تسک‌ها
    try:
        task_scheduler_instance = TaskScheduler(agent_runner_callback=run_agent)
        task_scheduler_instance.start()
    except Exception as e:
        print(f"Could not start TaskScheduler: {e}")

    print("================================")
    print("AI Assistant - WORKSPACE MODE")
    print("Model:", MODEL)
    print("http://127.0.0.1:5000")
    print("================================")
    app.run(host="127.0.0.1", port=5000, debug=False)